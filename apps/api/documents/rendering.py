"""Bounded local conversion of saved text. No network converters or embedded HTML."""

import io
import threading
import zipfile
import zlib
from datetime import UTC, datetime
from pathlib import Path

from identity.security import IdentityError

from documents.models import DocumentVersion

FONT_PATH = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
FONT_LOCK = threading.Lock()


def render_pdf(value: DocumentVersion) -> bytes:
    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas

    if not FONT_PATH.is_file():
        raise IdentityError(
            "EXPORT_UNAVAILABLE", 503, "PDF export unavailable; keep editing manually"
        )
    with FONT_LOCK:
        if "MinhaDocument" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("MinhaDocument", str(FONT_PATH)))
    stream = io.BytesIO()
    canvas = Canvas(stream, pagesize=A4, invariant=1, pageCompression=1)
    canvas.setTitle(value.title)
    width, height = A4
    font_size, margin, line_height = 11, 45, 17
    maximum = width - margin * 2
    y = height - margin

    def display(text: str) -> str:
        return get_display(arabic_reshaper.reshape(text)) if value.language == "ar" else text

    def draw(text: str) -> None:
        nonlocal y
        if y < margin:
            canvas.showPage()
            y = height - margin
        canvas.setFont("MinhaDocument", font_size)
        shaped = display(text)
        if value.language == "ar":
            canvas.drawRightString(width - margin, y, shaped)
        else:
            canvas.drawString(margin, y, shaped)
        y -= line_height

    # Character-based wrapping prevents arbitrarily long single words escaping margins.
    for paragraph in (value.title + "\n\n" + value.content).splitlines():
        line = ""
        for char in paragraph.replace("\t", "    "):
            candidate = line + char
            if (
                line
                and pdfmetrics.stringWidth(display(candidate), "MinhaDocument", font_size) > maximum
            ):
                draw(line)
                line = char
            else:
                line = candidate
        draw(line)
    canvas.save()
    result = stream.getvalue()
    if not result.startswith(b"%PDF-") or b"%%EOF" not in result[-64:]:
        raise IdentityError("EXPORT_INVALID", 503, "Export could not be verified")
    return result


def render_docx(value: DocumentVersion) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    output = Document()
    output.core_properties.title = value.title
    output.core_properties.created = datetime(2000, 1, 1, tzinfo=UTC)
    output.core_properties.modified = datetime(2000, 1, 1, tzinfo=UTC)
    output.core_properties.author = ""
    output.core_properties.last_modified_by = ""
    output.add_heading(value.title, 0)
    canonical_text = value.content.replace("\r\n", "\n").replace("\r", "\n")
    for text in canonical_text.split("\n"):
        paragraph = output.add_paragraph(text)
        if value.language == "ar":
            properties = paragraph._p.get_or_add_pPr()
            bidi = OxmlElement("w:bidi")
            bidi.set(qn("w:val"), "1")
            properties.append(bidi)
        for run in paragraph.runs:
            language = OxmlElement("w:lang")
            language.set(qn("w:val"), value.language)
            run._r.get_or_add_rPr().append(language)
    stream = io.BytesIO()
    output.save(stream)
    # python-docx uses wall-clock ZIP timestamps. Canonicalize the container so a
    # retry after accepted storage but before the export row saves identical bytes.
    canonical = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(canonical, "w") as target:
        for name in sorted(source.namelist()):
            payload = source.read(name)
            metadata = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            # Highly repetitive legitimate text must not trip the storage ZIP-bomb
            # ratio bound. Store only such entries uncompressed, keeping limits intact.
            compressed_size = max(1, len(zlib.compress(payload)) - 6)
            metadata.compress_type = (
                zipfile.ZIP_STORED if len(payload) > compressed_size * 80 else zipfile.ZIP_DEFLATED
            )
            metadata.create_system = 3
            metadata.external_attr = 0o600 << 16
            target.writestr(metadata, payload)
    result = canonical.getvalue()
    # Parse our own final container rather than treating a successful conversion as proof.
    parsed = Document(io.BytesIO(result))
    if "\n".join(item.text for item in parsed.paragraphs[1:]) != canonical_text:
        raise IdentityError("EXPORT_INVALID", 503, "Export could not be verified")
    return result


def render(value: DocumentVersion, output_format: str) -> tuple[bytes, str]:
    try:
        if output_format == "PDF":
            return render_pdf(value), "application/pdf"
        if output_format == "DOCX":
            return (
                render_docx(value),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
    except IdentityError:
        raise
    except Exception as exc:
        raise IdentityError(
            "EXPORT_UNAVAILABLE", 503, "Export unavailable; saved document is preserved"
        ) from exc
    raise IdentityError("INVALID_EXPORT_FORMAT", 400, "Choose PDF or DOCX")
