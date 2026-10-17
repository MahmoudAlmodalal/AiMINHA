"""Actual local rendering and mandatory actual scanner on private exports."""

import io
import uuid
from unittest.mock import patch

import pytest
from django.db import DatabaseError, transaction
from documents import rendering, services
from documents.models import DocumentExport, DocumentVersion
from docx import Document as WordDocument
from identity.models import IdentitySession
from identity.security import IdentityError, digest
from private_files import scanner
from private_files.models import ScanReceipt
from private_files.services import read_owned_bytes
from pypdf import PdfReader
from test_documents import create, mutate, signed_client
from test_documents import document_settings as document_settings

pytestmark = pytest.mark.django_db


@pytest.fixture
def export_setup(settings, tmp_path):
    settings.PRIVATE_FILES_ROOT = str(tmp_path / "private-exports")
    settings.PRIVATE_FILES_STORAGE_APPROVED = True
    client, user = signed_client()
    value = create(
        client,
        language="ar",
        title="سيرة ذاتية اصطناعية",
        content="اسم الطالب الاصطناعي\nدراسة الهندسة",
    )
    version = DocumentVersion.objects.get(pk=value["current"]["id"])
    proof = IdentitySession.objects.get(digest=digest(client.session.session_key))
    services.approve_version(
        user, proof, version.document_id, version.pk, version.content_hash, True
    )
    return client, user, proof, value, version


@pytest.fixture
def actual_scanner():
    try:
        return scanner.health()
    except IdentityError:
        pytest.skip("Actual ClamAV unavailable; actual export acceptance was not tested")


@pytest.mark.parametrize("output_format", ["PDF", "DOCX"])
def test_local_rendered_exports_parse_and_include_unicode(export_setup, output_format):
    _, _, _, _, version = export_setup
    data, mime = rendering.render(version, output_format)
    if output_format == "PDF":
        parsed = PdfReader(io.BytesIO(data))
        assert len(parsed.pages) >= 1 and "application/pdf" == mime
        assert parsed.pages[0].extract_text()
    else:
        parsed = WordDocument(io.BytesIO(data))
        assert version.content == "\n".join(item.text for item in parsed.paragraphs[1:])
        assert "wordprocessingml" in mime


@pytest.mark.parametrize("output_format", ["PDF", "DOCX"])
def test_scanner_failure_preserves_approved_artifact_retry_exports_once_without_credit(
    export_setup, actual_scanner, output_format
):
    from credits.models import Operation

    client, user, proof, value, version = export_setup
    before = Operation.objects.count()
    key = str(uuid.uuid4())
    path = f"documents/{value['id']}/exports"
    request = {"version_id": str(version.pk), "format": output_format}
    with patch("private_files.scanner.scan", side_effect=scanner.unavailable()):
        failed = mutate(client, path, request, key=key)
    assert failed.status_code == 503, failed.content
    assert not DocumentExport.objects.exists()
    version.refresh_from_db()
    assert version.content == value["current"]["content"]
    assert services.version_data(version)["approved"]
    result = mutate(client, path, request, key=key)
    assert result.status_code == 200, result.content
    duplicate = mutate(client, path, request, key=key)
    assert duplicate.status_code == 200 and duplicate.json() == result.json()
    assert DocumentExport.objects.count() == ScanReceipt.objects.count() == 1
    assert Operation.objects.count() == before
    saved = DocumentExport.objects.get()
    assert saved.version_id == version.pk and saved.content_hash == version.content_hash
    content = read_owned_bytes(user.pk, saved.asset_id)
    assert content.startswith(b"%PDF-") if output_format == "PDF" else content.startswith(b"PK")
    assert (
        services.export_version(user, proof, version.document_id, version.pk, output_format).pk
        == saved.pk
    )


def test_unapproved_new_version_cannot_export_and_foreign_version_is_denied(export_setup):
    client, user, proof, _, version = export_setup
    current = services.save_version(
        user,
        proof,
        version.document_id,
        {"title": "Edited", "language": "en", "content": "Changed manual facts"},
        1,
    )
    with pytest.raises(IdentityError, match="Approve"):
        services.export_version(user, proof, current.pk, current.current_id, "PDF")
    stranger, _ = signed_client()
    result = mutate(
        stranger,
        f"documents/{version.document_id}/exports",
        {"version_id": str(version.pk), "format": "PDF"},
    )
    assert result.status_code == 403 and not DocumentExport.objects.exists()


def test_converter_failure_keeps_saved_approval_and_pdf_missing_font_fails_closed(
    export_setup, tmp_path
):
    _, user, proof, _, version = export_setup
    with patch("documents.rendering.FONT_PATH", tmp_path / "missing.ttf"):
        with pytest.raises(IdentityError, match="PDF export unavailable"):
            services.export_version(user, proof, version.document_id, version.pk, "PDF")
    assert DocumentVersion.objects.get(pk=version.pk).content_hash == version.content_hash
    assert services.version_data(version)["approved"] and not DocumentExport.objects.exists()


def test_frozen_export_cannot_be_deleted_or_replaced(export_setup, actual_scanner):
    from private_files.services import delete_owned

    _, user, proof, _, version = export_setup
    saved = services.export_version(user, proof, version.document_id, version.pk, "DOCX")
    snapshot = services.freeze_approved_version(
        user, proof, version.document_id, version.pk, trusted_application_id=uuid.uuid4()
    )
    assert services.asset_is_frozen(saved.asset_id)
    with pytest.raises(IdentityError):
        delete_owned(user.pk, saved.asset_id)
    with pytest.raises(DatabaseError), transaction.atomic():
        DocumentExport.objects.filter(pk=saved.pk).update(content_hash="a" * 64)
    snapshot.refresh_from_db()
    assert snapshot.version_id == version.pk


def test_export_crash_after_scanned_storage_recovers_same_exact_bytes(export_setup, actual_scanner):
    from private_files.models import StoredFile

    _, user, proof, _, version = export_setup
    with patch("zipfile.time.localtime", return_value=(2020, 1, 1, 0, 0, 0, 2, 1, 0)):
        original_bytes, _ = rendering.render(version, "DOCX")
        with patch("documents.services.audit", side_effect=RuntimeError("Synthetic audit outage")):
            with pytest.raises(RuntimeError, match="Synthetic audit outage"):
                services.export_version(user, proof, version.document_id, version.pk, "DOCX")
    assert not DocumentExport.objects.exists()
    assert ScanReceipt.objects.count() == 1
    with patch("zipfile.time.localtime", return_value=(2030, 1, 1, 0, 0, 0, 2, 1, 0)):
        retry_bytes, _ = rendering.render(version, "DOCX")
        saved = services.export_version(user, proof, version.document_id, version.pk, "DOCX")
    assert retry_bytes == original_bytes
    assert read_owned_bytes(user.pk, saved.asset_id) == original_bytes
    assert DocumentExport.objects.count() == ScanReceipt.objects.count() == 1
    assert StoredFile.objects.filter(capacity_state="RETAINED").count() == 1


@pytest.mark.parametrize("content", ["Saved approved content\n", "One\rTwo", "One\r\nTwo"])
def test_docx_render_preserves_lines_and_trailing_empty_paragraph(export_setup, content):
    *_, version = export_setup
    version.content = content
    data, _ = rendering.render(version, "DOCX")
    parsed = WordDocument(io.BytesIO(data))
    expected = content.replace("\r\n", "\n").replace("\r", "\n")
    assert "\n".join(item.text for item in parsed.paragraphs[1:]) == expected


def test_export_request_key_binds_exact_version_format_before_io(export_setup):
    from documents.models import DocumentExportRequest

    _, user, proof, _, version = export_setup
    key = str(uuid.uuid4())
    services.bind_export_request(user, proof, version.document_id, version.pk, "PDF", key)
    services.bind_export_request(user, proof, version.document_id, version.pk, "PDF", key)
    assert DocumentExportRequest.objects.count() == 1
    with pytest.raises(IdentityError, match="Idempotency key already used"):
        services.bind_export_request(user, proof, version.document_id, version.pk, "DOCX", key)
    intent = DocumentExportRequest.objects.get()
    assert intent.key_digest != key
    with pytest.raises(DatabaseError), transaction.atomic():
        DocumentExportRequest.objects.filter(pk=intent.pk).update(format="DOCX")


@pytest.mark.parametrize("content", ["a" * 40000, "ن" * 40000])
def test_long_repetitive_valid_docx_stays_within_container_limits(export_setup, content):
    from private_files.validation import envelope, structural

    *_, version = export_setup
    version.content = content
    data, mime = rendering.render(version, "DOCX")
    assert len(data) < 5_000_000
    envelope("document.docx", mime, data)
    structural(mime, data)
    parsed = WordDocument(io.BytesIO(data))
    assert parsed.paragraphs[1].text == content


def test_oversized_export_is_rejected_without_losing_saved_document(export_setup):
    _, user, proof, _, version = export_setup
    oversized = b"%PDF-" + b"x" * 5_000_000
    with patch("documents.rendering.render", return_value=(oversized, "application/pdf")):
        with pytest.raises(IdentityError) as error:
            services.export_version(user, proof, version.document_id, version.pk, "PDF")
    assert error.value.code == "FILE_TOO_LARGE"
    assert not DocumentExport.objects.exists() and not ScanReceipt.objects.exists()
    version.refresh_from_db()
    assert services.version_data(version)["approved"]
