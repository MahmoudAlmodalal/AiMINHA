import io
import os
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from threading import Barrier
from unittest.mock import patch

import pytest
from django.db import IntegrityError, close_old_connections, connection, transaction
from django.utils import timezone
from identity.models import Account, IdentitySession
from identity.security import IdentityError, digest
from PIL import Image
from privacy.models import PrivacyAudit, PrivateAsset, TicketAccess, TicketShare
from private_files import scanner, services, storage, validation
from private_files.models import ScanReceipt, StorageLock, StoredFile
from rest_framework.test import APIClient

pytestmark = pytest.mark.django_db


def student(role="STUDENT"):
    return Account.objects.create_user(
        str(uuid.uuid4()) + "@example.invalid",
        "synthetic-only-passphrase-529!",
        birth_date=date(2000, 1, 1),
        locale="en",
        state="ACTIVE",
        role=role,
        email_verified_at=timezone.now(),
        admin_scopes=["SUPPORT"],
    )


def png():
    buffer = io.BytesIO()
    Image.new("RGB", (3, 3), "blue").save(buffer, format="PNG")
    return buffer.getvalue()


def proof(actor):
    return IdentitySession.objects.create(
        account=actor,
        epoch=actor.session_epoch,
        digest=digest(str(uuid.uuid4())),
        expires_at=timezone.now() + timedelta(days=1),
        mfa_verified_at=timezone.now(),
    )


@pytest.fixture(autouse=True)
def private_settings(settings, tmp_path):
    StorageLock.objects.get_or_create(pk=1)
    settings.PRIVATE_FILES_ROOT = str(tmp_path / "private")
    settings.PRIVATE_FILES_STORAGE_APPROVED = True
    settings.PASSWORD_HASHERS = ["django.contrib.auth.hashers.MD5PasswordHasher"]
    settings.ALLOWED_HOSTS = ["testserver"]
    clean = scanner.CleanScan("ClamAV synthetic-test", "12345", timezone.now())
    with (
        patch("private_files.scanner.health", return_value=clean),
        patch(
            "private_files.scanner.scan",
            return_value=clean,
        ),
    ):
        yield


def save(actor=None, key=None):
    actor = actor or student()
    asset = services.store_bytes(actor.pk, "cv.png", "image/png", png(), key or str(uuid.uuid4()))
    return actor, asset


def code(expected, action):
    with pytest.raises(IdentityError) as error:
        action()
    assert error.value.code == expected
    return error.value


def authenticated(actor):
    client = APIClient(enforce_csrf_checks=True)
    client.force_login(actor)
    IdentitySession.objects.create(
        account=actor,
        digest=digest(client.session.session_key),
        epoch=actor.session_epoch,
        expires_at=timezone.now() + timedelta(days=1),
        mfa_verified_at=timezone.now(),
    )
    return client


def post(client, path, data=None, **kwargs):
    token = client.get("/api/v1/session/csrf").json()["csrf_token"]
    return client.post(path, data or {}, HTTP_X_CSRFTOKEN=token, **kwargs)


def test_real_accepted_byte_identity_receipt_and_owner_read():
    actor, asset = save()
    file = StoredFile.objects.get(pk=asset.pk)
    receipt = ScanReceipt.objects.get(file=file)
    assert asset.state == "ACCEPTED"
    assert receipt.sha256 == file.sha256 == asset.object_version
    assert file.storage_token == asset.object_key
    assert receipt.size == file.size == asset.declared_size
    assert services.read_owned_bytes(actor.pk, asset.pk) == png()
    assert file.capacity_bytes == len(png()) and file.capacity_state == "RETAINED"
    assert not hasattr(asset, "url")


def test_metadata_only_old_fixture_can_never_yield_bytes():
    actor = student()
    metadata = PrivateAsset.objects.create(owner=actor, state="ACCEPTED", object_version="legacy")
    code("OWNERSHIP_REQUIRED", lambda: services.read_owned_bytes(actor.pk, metadata.pk))


@pytest.mark.parametrize("role", ["STUDENT", "ADMIN"])
def test_cross_owner_or_staff_bytes_denied_and_audited(role):
    actor, asset = save()
    other = student(role)
    client = authenticated(other)
    response = post(client, f"/api/v1/files/{asset.pk}/access-grants", format="json")
    assert response.status_code == (403 if role == "ADMIN" else 404)
    assert PrivacyAudit.objects.filter(actor_id=other.pk, outcome="DENY").exists()
    assert str(actor.pk) not in response.content.decode()


@pytest.mark.parametrize(
    "name,mime,data",
    [
        ("cv.pdf", "application/pdf", b"plain not a PDF"),
        ("cv.png", "application/pdf", png()),
        ("cv.exe", "image/png", png()),
        (
            "cv.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            png(),
        ),
        ("../cv.png", "image/png", png()),
        ("cv.png", "image/png; charset=utf-8", png()),
    ],
)
def test_mime_extension_signature_spoof_rejected_without_storage(name, mime, data):
    code("FILE_TYPE_REJECTED", lambda: services.store_bytes(student().pk, name, mime, data, "key"))
    assert not StoredFile.objects.exists()


@pytest.mark.parametrize("data", [b"", b"x" * 5_000_001])
def test_oversized_and_empty_bytes_rejected(data):
    code(
        "FILE_TOO_LARGE",
        lambda: services.store_bytes(student().pk, "cv.png", "image/png", data, "key"),
    )
    assert not StoredFile.objects.exists()


def test_size_boundary_declared_cap_actual_measured_not_client_claim():
    actor = student()
    with patch("private_files.validation.structural"):
        asset = services.store_bytes(
            actor.pk, "cv.png", "image/png", png() + b"x" * (5_000_000 - len(png())), "key"
        )
    assert asset.declared_size == 5_000_000


def test_duplicate_upload_same_bytes_once_and_mismatch_conflicts():
    actor, asset = save(key="repeat")
    assert services.store_bytes(actor.pk, "cv.png", "image/png", png(), "repeat").pk == asset.pk
    assert StoredFile.objects.count() == 1
    code(
        "IDEMPOTENCY_CONFLICT",
        lambda: services.store_bytes(actor.pk, "different.png", "image/png", png(), "repeat"),
    )


def test_scanner_unavailable_before_reservation_offers_manual_entry():
    with patch("private_files.scanner.health", side_effect=scanner.unavailable()):
        error = code("FILE_SCANNER_UNAVAILABLE", lambda: save())
        assert "manual" in error.safe_title
        assert services.capabilities()["manual_entry_available"]
        assert not services.capabilities()["upload_available"]
    assert not StoredFile.objects.exists()


def test_scanner_error_after_write_purges_quota_then_retry_same_key():
    actor = student()
    with patch("private_files.scanner.scan", side_effect=scanner.unavailable()):
        code("FILE_SCANNER_UNAVAILABLE", lambda: save(actor, "recovery"))
    first = StoredFile.objects.get()
    assert first.capacity_bytes == 0 and first.capacity_state == "RELEASED"
    assert first.asset.state == "PURGED" and not ScanReceipt.objects.exists()
    _, recovered = save(actor, "recovery")
    assert recovered.state == "ACCEPTED" and recovered.pk != first.pk


def test_infected_rejected_no_clean_receipt_or_access():
    with patch(
        "private_files.scanner.scan", side_effect=IdentityError("FILE_MALWARE_DETECTED", 422)
    ):
        code("FILE_MALWARE_DETECTED", lambda: save())
    assert not PrivateAsset.objects.filter(state="ACCEPTED").exists()
    assert not ScanReceipt.objects.exists()


def test_cleanup_unknown_retains_capacity_and_blocks_duplicate_operation():
    actor = student()
    with (
        patch("private_files.scanner.scan", side_effect=scanner.unavailable()),
        patch(
            "private_files.storage.remove",
            side_effect=OSError,
        ),
    ):
        code("FILE_SCANNER_UNAVAILABLE", lambda: save(actor, "uncertain"))
    file = StoredFile.objects.get()
    assert file.capacity_bytes == len(png()) and file.asset.state == "PURGE_PENDING"
    code("FILE_NOT_READY", lambda: save(actor, "uncertain"))
    assert StoredFile.objects.count() == 1
    services.delete_owned(actor.pk, file.pk)
    file.refresh_from_db()
    assert file.capacity_bytes == 0


@pytest.mark.parametrize("global_cap", [False, True])
def test_owner_global_quota_allocation_and_once_release(global_cap):
    actor, asset = save()
    _, second = save(actor)
    bound = "MAX_GLOBAL_BYTES" if global_cap else "MAX_STUDENT_BYTES"
    with patch.object(validation, bound, len(png()) * 2):
        code("STORAGE_QUOTA_EXCEEDED", lambda: save(student() if global_cap else actor))
        services.delete_owned(actor.pk, asset.pk)
        services.delete_owned(actor.pk, asset.pk)
        _, third = save(actor)
    assert third.state == "ACCEPTED" and second.state == "ACCEPTED"


def test_scan_success_corrupt_content_never_accepted():
    code(
        "FILE_TYPE_REJECTED",
        lambda: services.store_bytes(
            student().pk, "bad.png", "image/png", b"\x89PNG\r\n\x1a\nbroken", "bad"
        ),
    )
    assert not PrivateAsset.objects.filter(state="ACCEPTED").exists()


def test_tamper_storage_after_scan_denies_read():
    actor, asset = save()
    file = StoredFile.objects.get(pk=asset.pk)
    path = storage.root() / (str(file.storage_token) + ".blob")
    path.write_bytes(png() + b"changed")
    code("FILE_INTEGRITY_FAILED", lambda: services.read_owned_bytes(actor.pk, asset.pk))


def test_non_overwriting_storage_and_symlink_denied(tmp_path):
    actor, asset = save()
    file = StoredFile.objects.get(pk=asset.pk)
    with pytest.raises(FileExistsError):
        storage.write(file, png())
    path = storage.root() / (str(file.storage_token) + ".blob")
    path.unlink()
    target = tmp_path / "unrelated"
    target.write_bytes(png())
    path.symlink_to(target)
    code("PRIVATE_STORAGE_UNAVAILABLE", lambda: services.read_owned_bytes(actor.pk, asset.pk))


def test_expiring_grant_requires_same_actor_session_and_asset():
    actor, asset = save()
    session = proof(actor)
    grant = services.issue_grant(actor, session, asset.pk)
    assert services.read_granted(actor, session, asset.pk, grant["token"])[1] == png()
    other = student()
    code(
        "ACCESS_GRANT_INVALID",
        lambda: services.read_granted(other, proof(other), asset.pk, grant["token"]),
    )
    code(
        "ACCESS_GRANT_INVALID",
        lambda: services.read_granted(actor, proof(actor), asset.pk, grant["token"]),
    )
    with patch(
        "private_files.services.timezone.now", return_value=timezone.now() + timedelta(seconds=61)
    ):
        code(
            "ACCESS_GRANT_INVALID",
            lambda: services.read_granted(actor, session, asset.pk, grant["token"]),
        )


@pytest.mark.parametrize("revocation", ["session", "account", "asset"])
def test_revocation_after_grant_denies_byte_delivery(revocation):
    actor, asset = save()
    session = proof(actor)
    grant = services.issue_grant(actor, session, asset.pk)
    if revocation == "session":
        IdentitySession.objects.filter(pk=session.pk).update(revoked_at=timezone.now())
    elif revocation == "account":
        Account.objects.filter(pk=actor.pk).update(state="DELETION_PENDING")
    else:
        services.delete_owned(actor.pk, asset.pk)
    with pytest.raises(IdentityError):
        services.read_granted(actor, session, asset.pk, grant["token"])


def test_account_revocation_during_storage_io_rechecked():
    actor, asset = save()
    session = proof(actor)
    grant = services.issue_grant(actor, session, asset.pk)
    original = storage.read

    def revoke(file):
        data = original(file)
        Account.objects.filter(pk=actor.pk).update(state="DELETION_PENDING")
        return data

    with patch("private_files.storage.read", side_effect=revoke):
        code(
            "VERIFIED_ADULT_REQUIRED",
            lambda: services.read_granted(actor, session, asset.pk, grant["token"]),
        )


@pytest.mark.parametrize("change", ["share", "assignment", "mfa", "scope", "owner"])
def test_support_only_explicit_assigned_attachment_and_current_authorization(change):
    owner, asset = save()
    admin = student("ADMIN")
    session = proof(admin)
    ticket = TicketAccess.objects.create(owner=owner, assigned_admin=admin)
    share = TicketShare.objects.create(owner=owner, asset=asset, ticket=ticket)
    grant = services.issue_grant(admin, session, asset.pk, share)
    assert services.read_granted(admin, session, asset.pk, grant["token"], share)[1] == png()
    if change == "share":
        TicketShare.objects.filter(pk=share.pk).update(revoked_at=timezone.now())
    elif change == "assignment":
        TicketAccess.objects.filter(pk=ticket.pk).update(assigned_admin=None)
    elif change == "mfa":
        IdentitySession.objects.filter(pk=session.pk).update(mfa_verified_at=None)
    elif change == "scope":
        Account.objects.filter(pk=admin.pk).update(admin_scopes=[])
    else:
        Account.objects.filter(pk=owner.pk).update(state="DELETION_PENDING")
    with pytest.raises(IdentityError):
        services.read_granted(admin, session, asset.pk, grant["token"], share)
    assert PrivacyAudit.objects.filter(actor_id=admin.pk, action="ticket_file_bytes_read").exists()


def test_upload_http_and_private_download_require_csrf_and_header():
    actor = student()
    client = authenticated(actor)
    file = io.BytesIO(png())
    file.name = "cv.png"
    from django.core.files.uploadedfile import SimpleUploadedFile

    response = post(
        client,
        "/api/v1/files/uploads",
        {"file": SimpleUploadedFile("cv.png", png(), content_type="image/png")},
        format="multipart",
        HTTP_IDEMPOTENCY_KEY="http-upload",
    )
    assert response.status_code == 201, response.content
    asset_id = response.json()["id"]
    no_grant = client.get(f"/api/v1/files/{asset_id}/content")
    assert no_grant.status_code == 403
    response = post(client, f"/api/v1/files/{asset_id}/access-grants", format="json")
    assert response.status_code == 201, response.content
    response = client.get(
        f"/api/v1/files/{asset_id}/content", HTTP_X_FILE_ACCESS_GRANT=response.json()["token"]
    )
    assert response.status_code == 200 and response.content == png()
    assert response["Content-Disposition"].endswith('private-file.png"')
    assert (
        "no-store" in response["Cache-Control"] and response["X-Content-Type-Options"] == "nosniff"
    )
    assert (
        client.post(f"/api/v1/files/{asset_id}/access-grants", {}, format="json").status_code == 403
    )


def test_http_oversize_cannot_create_temp_or_stored_object():
    actor = student()
    client = authenticated(actor)
    from django.core.files.uploadedfile import SimpleUploadedFile

    response = post(
        client,
        "/api/v1/files/uploads",
        {"file": SimpleUploadedFile("cv.png", png() + b"x" * 5_000_001, content_type="image/png")},
        format="multipart",
        HTTP_IDEMPOTENCY_KEY="oversized",
    )
    assert response.status_code == 413, response.content
    assert not StoredFile.objects.exists()


def test_acceptance_sql_guard_rejects_missing_scan_receipt():
    actor = student()
    with pytest.raises(IntegrityError), transaction.atomic():
        asset = PrivateAsset.objects.create(
            owner=actor, object_version="a" * 64, declared_size=1, state="ACCEPTED"
        )
        StoredFile.objects.create(
            asset=asset,
            storage_token=asset.object_key,
            operation_digest="a" * 64,
            sha256="a" * 64,
            size=1,
            mime="image/png",
            name="a.png",
            capacity_state="RETAINED",
            capacity_bytes=1,
        )
        with connection.cursor() as cursor:
            cursor.execute("SET CONSTRAINTS files_asset_exact IMMEDIATE")


def test_scan_receipt_and_stored_identity_are_database_immutable():
    _, asset = save()
    for model, values in [(StoredFile, {"sha256": "b" * 64}), (ScanReceipt, {"verdict": "FAILED"})]:
        with pytest.raises(IntegrityError), transaction.atomic():
            model.objects.filter(pk=asset.pk).update(**values)


@pytest.mark.parametrize("after_write", [False, True])
def test_abandoned_upload_crash_recovery_fences_late_writer_and_retries(after_write):
    actor = student()
    write = storage.write

    def crash(file, data):
        if after_write:
            write(file, data)
        raise SystemExit("synthetic process crash")

    with patch("private_files.storage.write", side_effect=crash), pytest.raises(SystemExit):
        save(actor, "crash-retry")
    file = StoredFile.objects.select_related("asset").get()
    assert file.capacity_state == "RESERVED" and file.asset.state == "QUARANTINED"
    assert file.capacity_bytes == len(png())
    assert services.recover_abandoned() == {"recovered": 0, "pending": 0}
    with patch(
        "private_files.services.timezone.now", return_value=timezone.now() + timedelta(minutes=6)
    ):
        assert services.recover_abandoned() == {"recovered": 1, "pending": 0}
    file.refresh_from_db()
    assert file.capacity_bytes == 0 and file.capacity_state == "RELEASED"
    assert not (storage.root() / (str(file.storage_token) + ".blob")).exists()
    with storage.lease(file):
        code("FILE_LEASE_EXPIRED", lambda: services._assert_writable(file))
    assert not (storage.root() / (str(file.storage_token) + ".blob")).exists()
    _, retry = save(actor, "crash-retry")
    assert retry.state == "ACCEPTED" and retry.pk != file.pk


def test_recovery_skips_active_writer_and_retains_uncertain_cleanup_quota():
    actor = student()
    with patch("private_files.storage.write", side_effect=SystemExit), pytest.raises(SystemExit):
        save(actor, "active-writer")
    file = StoredFile.objects.get()
    future = timezone.now() + timedelta(minutes=6)
    with storage.lease(file), patch("private_files.services.timezone.now", return_value=future):
        assert services.recover_abandoned() == {"recovered": 0, "pending": 1}
    file.refresh_from_db()
    assert file.capacity_state == "RESERVED"
    with (
        patch("private_files.services.timezone.now", return_value=future),
        patch(
            "private_files.storage.remove",
            side_effect=OSError,
        ),
    ):
        assert services.recover_abandoned() == {"recovered": 0, "pending": 1}
    file.refresh_from_db()
    assert file.capacity_bytes == len(png())
    with patch("private_files.services.timezone.now", return_value=future):
        assert services.recover_abandoned() == {"recovered": 1, "pending": 0}


def test_slow_scanner_past_lease_never_accepts_or_leaves_bytes():
    initial = timezone.now()
    clock = {"now": initial}
    clean = scanner.CleanScan("ClamAV synthetic-test", "12345", initial)

    def slow(data):
        clock["now"] = initial + timedelta(minutes=6)
        return clean

    with (
        patch("private_files.scanner.scan", side_effect=slow),
        patch(
            "private_files.services.timezone.now",
            side_effect=lambda: clock["now"],
        ),
    ):
        code("FILE_NOT_READY", lambda: save())
    assert not ScanReceipt.objects.exists()
    assert StoredFile.objects.get().capacity_bytes == 0


@pytest.mark.django_db(transaction=True)
@pytest.mark.parametrize("bound", ["MAX_STUDENT_BYTES", "MAX_GLOBAL_BYTES"])
def test_real_postgres_concurrent_quota_admission_counts_reserved_bytes(bound):
    actor = student()
    second = actor if bound == "MAX_STUDENT_BYTES" else student()
    barrier = Barrier(2)
    clean = scanner.CleanScan("ClamAV synthetic-test", "12345", timezone.now())

    def ready():
        barrier.wait(timeout=5)
        return clean

    def run(owner):
        close_old_connections()
        try:
            save(owner)
            return "ACCEPTED"
        except IdentityError as error:
            return error.code
        finally:
            close_old_connections()

    with (
        patch.object(validation, bound, len(png())),
        patch(
            "private_files.scanner.health",
            side_effect=ready,
        ),
        ThreadPoolExecutor(max_workers=2) as pool,
    ):
        outcomes = list(pool.map(run, [actor, second]))
    assert sorted(outcomes) == ["ACCEPTED", "STORAGE_QUOTA_EXCEEDED"]
    assert StoredFile.objects.count() == 1 and ScanReceipt.objects.count() == 1


@pytest.mark.parametrize("location", ["apps/web/public/private", "apps/web/.next/static/private"])
def test_public_storage_misconfiguration_rejected(settings, location):
    settings.PRIVATE_FILES_ROOT = str(settings.ROOT / location)
    code("PRIVATE_STORAGE_UNAVAILABLE", storage.root)


def test_symlink_storage_ancestor_rejected(settings, tmp_path):
    actual = tmp_path / "actual"
    actual.mkdir()
    link = tmp_path / "alias"
    link.symlink_to(actual, target_is_directory=True)
    settings.PRIVATE_FILES_ROOT = str(link / "private")
    code("PRIVATE_STORAGE_UNAVAILABLE", storage.root)


def test_zip_with_docx_names_but_false_opc_roots_is_not_a_word_document():
    data = docx(b"<foo/>")
    code(
        "FILE_TYPE_REJECTED",
        lambda: services.store_bytes(
            student().pk, "spoof.docx", validation.MIMES[".docx"], data, "spoof"
        ),
    )
    assert not PrivateAsset.objects.filter(state="ACCEPTED").exists()


def test_valid_docx_independent_parser_round_trip_accepted():
    from docx import Document

    document = Document()
    document.add_paragraph("Synthetic confirmed CV facts")
    output = io.BytesIO()
    document.save(output)
    actor = student()
    asset = services.store_bytes(
        actor.pk, "real.docx", validation.MIMES[".docx"], output.getvalue(), "real-docx"
    )
    parsed = Document(io.BytesIO(services.read_owned_bytes(actor.pk, asset.pk)))
    assert parsed.paragraphs[0].text == "Synthetic confirmed CV facts"


def docx(content=None, extra=None):
    content = content or (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b"<w:body><w:p><w:r><w:t>Synthetic CV</w:t></w:r></w:p></w:body></w:document>"
    )
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            b'<Override PartName="/word/document.xml" '
            b'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            b"</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b'<Relationship Id="rId1" '
            b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
            b'officeDocument" '
            b'Target="word/document.xml"/></Relationships>',
        )
        archive.writestr("word/document.xml", content)
        if extra:
            archive.writestr(*extra)
    return data.getvalue()


@pytest.mark.parametrize(
    "data",
    [
        docx(b'<!DOCTYPE x [<!ENTITY y "z">]><document>&y;</document>'),
        docx(extra=("../secret.xml", b"<x/>")),
        docx(extra=("word/vbaProject.bin", b"macro")),
        docx(extra=("word/bomb.xml", b"x" * 200_000)),
        docx(
            extra=(
                "word/_rels/document.xml.rels",
                b'<Relationships><Relationship TargetMode="External" Target="http://example.invalid"/></Relationships>',
            )
        ),
    ],
)
def test_docx_xxe_traversal_macro_bomb_external_content_rejected(data):
    code(
        "FILE_TYPE_REJECTED",
        lambda: services.store_bytes(
            student().pk, "cv.docx", validation.MIMES[".docx"], data, "docx"
        ),
    )
    assert not PrivateAsset.objects.filter(state="ACCEPTED").exists()


@pytest.mark.skipif(
    not os.getenv("MINHAAI_REAL_SCANNER_TEST"),
    reason="Requires actual local ClamAV with current signatures",
)
def test_actual_clamav_clean_infected_outage_stale_signatures(settings):
    # Unpatch both callables: this opt-in integration proves the real adapter.
    from importlib import reload

    live = reload(scanner)
    assert live.health().engine.startswith("ClamAV ")
    assert live.scan(png()).signature_version.isdigit()
    eicar = b"X5O!P%@AP[4\\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*"
    code("FILE_MALWARE_DETECTED", lambda: live.scan(eicar))
    with patch.object(settings, "PRIVATE_FILES_SCANNER_PORT", 1):
        code("FILE_SCANNER_UNAVAILABLE", live.health)
    with patch.object(settings, "PRIVATE_FILES_SCANNER_MAX_SIGNATURE_AGE_HOURS", 0):
        code("FILE_SCANNER_UNAVAILABLE", live.health)
